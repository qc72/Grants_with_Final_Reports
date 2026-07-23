from __future__ import annotations

import io
import re
import shutil
import tempfile
import zipfile
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Iterable

import fitz  # PyMuPDF
from docx import Document

PROJECT_ID_RE = re.compile(r"\b([A-Za-z]{2,})\s+(\d{2}\.\d{3})\b")


@dataclass
class PdfCandidate:
    path: str
    filename: str
    score: int
    reasons: list[str] = field(default_factory=list)


@dataclass
class ProjectRecord:
    project_id: str
    folder_name: str
    title: str = ""
    program: str = ""
    academic_year: str = ""
    funding_amount: str = ""
    principal_investigator: str = ""
    pi_college: str = ""
    community_partners: str = ""
    student_involvement: str = ""
    number_of_students: str = ""
    community_need: str = ""
    community_impact: str = ""
    publications: str = ""
    cel_classification: str = ""
    confidence: str = ""
    brief_explanation: str = ""
    final_report_available: str = ""
    final_report_path: str = ""
    final_report_score: int = 0
    final_report_confidence: str = "Low"
    summary_text: str = ""
    highlight_text: str = ""
    project_note_text: str = ""
    missing_expected_files: list[str] = field(default_factory=list)
    files: list[str] = field(default_factory=list)
    pdf_candidates: list[PdfCandidate] = field(default_factory=list)

    def to_dict(self) -> dict:
        data = asdict(self)
        data["pdf_candidates"] = [asdict(x) for x in self.pdf_candidates]
        return data


def safe_extract_zip(zip_bytes: bytes, destination: Path) -> None:
    """Extract a ZIP while blocking traversal and obvious ZIP bombs."""
    destination = destination.resolve()
    with zipfile.ZipFile(io.BytesIO(zip_bytes)) as zf:
        members = zf.infolist()
        if len(members) > 50_000:
            raise ValueError("ZIP contains too many entries.")
        if sum(m.file_size for m in members) > 10 * 1024**3:
            raise ValueError("ZIP expands beyond the 10 GB safety limit.")
        for member in members:
            member_path = (destination / member.filename).resolve()
            if destination != member_path and destination not in member_path.parents:
                raise ValueError(f"Unsafe ZIP entry blocked: {member.filename}")
        zf.extractall(destination)


def read_text(path: Path) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    return path.read_text(errors="replace")


def read_docx(path: Path) -> str:
    doc = Document(path)
    chunks: list[str] = []
    chunks.extend(p.text.strip() for p in doc.paragraphs if p.text.strip())
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                chunks.append(" | ".join(cells))
    return "\n".join(chunks)


def read_pdf(path: Path, max_pages: int | None = None) -> str:
    chunks: list[str] = []
    with fitz.open(path) as doc:
        page_count = len(doc) if max_pages is None else min(len(doc), max_pages)
        for i in range(page_count):
            chunks.append(doc[i].get_text("text"))
    return "\n".join(chunks)


def normalize_label(label: str) -> str:
    label = label.strip().lower()
    label = re.sub(r"[^a-z0-9]+", " ", label)
    return re.sub(r"\s+", " ", label).strip()


def parse_key_values(text: str) -> dict[str, str]:
    values: dict[str, str] = {}
    for raw_line in text.splitlines():
        line = raw_line.strip().lstrip("•-")
        if not line or ":" not in line:
            continue
        key, value = line.split(":", 1)
        key = normalize_label(key)
        value = value.strip()
        if key and value:
            values[key] = value
    return values


ALIASES: dict[str, tuple[str, ...]] = {
    "title": ("title",),
    "program": ("program", "category program", "category"),
    "academic_year": ("academic year",),
    "funding_amount": ("funding amount", "total funding"),
    "principal_investigator": ("pi project lead", "principal investigator"),
    "pi_college": ("pi college or department", "pi college"),
    "community_partners": ("community partner s", "community partners"),
    "student_involvement": ("student involvement",),
    "number_of_students": ("number of students involved",),
    "community_need": ("main community need",),
    "community_impact": ("main community impact",),
    "publications": ("publications final products", "publications or final products"),
    "cel_classification": ("overall cel classification",),
    "confidence": ("confidence",),
    "brief_explanation": ("brief explanation",),
    "final_report_available": ("final report available",),
}


def first_value(*maps: dict[str, str], aliases: Iterable[str]) -> str:
    for mapping in maps:
        for alias in aliases:
            if alias in mapping and mapping[alias]:
                return mapping[alias]
    return ""


def score_final_report(path: Path, project_id: str) -> PdfCandidate:
    name = path.name.lower().replace("_", " ").replace("-", " ")
    score = 0
    reasons: list[str] = []

    def add(points: int, reason: str) -> None:
        nonlocal score
        score += points
        reasons.append(f"{points:+d}: {reason}")

    if project_id.lower() in name:
        add(3, "filename contains project ID")
    if "final report" in name:
        add(10, "filename says final report")
    elif "report" in name:
        add(3, "filename contains report")
    if "budget" in name:
        add(-10, "filename indicates budget")
    if "application" in name or "proposal" in name:
        add(-9, "filename indicates application/proposal")
    if "worksheet" in name:
        add(-7, "filename indicates worksheet")
    if "interim" in name or "progress" in name:
        add(-4, "filename indicates non-final report")

    try:
        text = read_pdf(path, max_pages=6).lower()
    except Exception as exc:  # corrupted/encrypted PDFs remain visible for manual review
        add(-20, f"PDF text could not be read ({type(exc).__name__})")
        return PdfCandidate(str(path), path.name, score, reasons)

    content_signals = [
        ("final report", 12, "document text says final report"),
        ("submitted on", 4, "contains submission metadata"),
        ("how have undergraduate students contributed", 7, "contains final-report questionnaire"),
        ("what has been the role", 4, "contains partner-role question"),
        ("overall, how satisfied", 5, "contains recipient satisfaction question"),
        ("would you recommend this funding", 4, "contains closing report question"),
    ]
    for needle, points, reason in content_signals:
        if needle in text:
            add(points, reason)

    if "project budget" in text and "final report" not in text:
        add(-8, "content appears to be a budget")
    if "application" in text[:2500] and "final report" not in text:
        add(-5, "opening content appears to be an application")

    return PdfCandidate(str(path), path.name, score, reasons)


def confidence_from_candidates(candidates: list[PdfCandidate]) -> str:
    if not candidates:
        return "None"
    ordered = sorted(candidates, key=lambda x: x.score, reverse=True)
    top = ordered[0].score
    gap = top - ordered[1].score if len(ordered) > 1 else top
    if top >= 18 and gap >= 8:
        return "High"
    if top >= 10 and gap >= 3:
        return "Medium"
    return "Low"


def identify_project_directories(root: Path) -> list[Path]:
    dirs: set[Path] = set()
    for path in root.rglob("*"):
        if not path.is_dir():
            continue
        if PROJECT_ID_RE.search(path.name):
            dirs.add(path)
    # Also detect folders containing a highlight or project note with an embedded ID.
    for path in root.rglob("*"):
        if not path.is_file() or path.suffix.lower() != ".txt":
            continue
        lower = path.name.lower()
        if "highlight" not in lower and not ("project" in lower and "note" in lower):
            continue
        try:
            text = read_text(path)[:2000]
        except OSError:
            continue
        if PROJECT_ID_RE.search(text):
            dirs.add(path.parent)
    return sorted(dirs, key=lambda p: str(p).lower())


def find_best_file(files: list[Path], *, suffix: str, include: tuple[str, ...]) -> Path | None:
    candidates = [
        p for p in files
        if p.suffix.lower() == suffix and all(token in p.name.lower() for token in include)
    ]
    if not candidates:
        return None
    return sorted(candidates, key=lambda p: (len(p.name), p.name.lower()))[0]


def project_id_from_folder_or_text(folder: Path, texts: Iterable[str]) -> str:
    match = PROJECT_ID_RE.search(folder.name)
    if match:
        return f"{match.group(1).upper()} {match.group(2)}"
    for text in texts:
        match = PROJECT_ID_RE.search(text)
        if match:
            return f"{match.group(1).upper()} {match.group(2)}"
    return folder.name


def scan_project(folder: Path) -> ProjectRecord:
    files = [p for p in folder.rglob("*") if p.is_file()]
    summary_path = find_best_file(files, suffix=".docx", include=("summary",))
    highlight_path = find_best_file(files, suffix=".txt", include=("highlight",))
    note_path = next(
        (p for p in files if p.suffix.lower() == ".txt" and "project" in p.name.lower() and "note" in p.name.lower()),
        None,
    )

    summary_text = read_docx(summary_path) if summary_path else ""
    highlight_text = read_text(highlight_path) if highlight_path else ""
    note_text = read_text(note_path) if note_path else ""
    project_id = project_id_from_folder_or_text(folder, (highlight_text, note_text, summary_text))

    highlight_kv = parse_key_values(highlight_text)
    note_kv = parse_key_values(note_text)
    summary_kv = parse_key_values(summary_text)

    pdfs = [p for p in files if p.suffix.lower() == ".pdf"]
    candidates = sorted(
        (score_final_report(p, project_id) for p in pdfs),
        key=lambda x: (x.score, x.filename.lower()),
        reverse=True,
    )
    selected = candidates[0] if candidates else None

    missing: list[str] = []
    if not summary_path:
        missing.append("summary.docx")
    if not highlight_path:
        missing.append("highlight.txt")
    if not note_path:
        missing.append("project_note.txt")
    if not selected or selected.score < 10:
        missing.append("confident final-report PDF")

    maps = (highlight_kv, note_kv, summary_kv)
    kwargs = {
        field_name: first_value(*maps, aliases=aliases)
        for field_name, aliases in ALIASES.items()
    }

    return ProjectRecord(
        project_id=project_id,
        folder_name=folder.name,
        **kwargs,
        final_report_path=selected.path if selected else "",
        final_report_score=selected.score if selected else 0,
        final_report_confidence=confidence_from_candidates(candidates),
        summary_text=summary_text,
        highlight_text=highlight_text,
        project_note_text=note_text,
        missing_expected_files=missing,
        files=[str(p) for p in sorted(files, key=lambda x: x.name.lower())],
        pdf_candidates=candidates,
    )


def scan_zip_bytes(zip_bytes: bytes) -> tuple[list[ProjectRecord], str]:
    temp_dir = Path(tempfile.mkdtemp(prefix="grant_insights_"))
    safe_extract_zip(zip_bytes, temp_dir)
    project_dirs = identify_project_directories(temp_dir)
    records = [scan_project(folder) for folder in project_dirs]
    return records, str(temp_dir)


def remove_scan_directory(path: str) -> None:
    shutil.rmtree(path, ignore_errors=True)
